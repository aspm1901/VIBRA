"""
Full implementation to update Chapter 6 in VIBRA.docx with bulletproof mathematical precision,
refined operational sequence, multi-criteria tie-breakers, explicit guardrails, and optimal pagination.
Includes all peer-review refinements:
1. Taxonomy explicitly identifies chalecos in 'Casacas, blazers y chalecos'.
2. Universal application of Paso 1 hard filters across all 3 search modes.
3. Natural distribution of 'Por estilo' across all 3 bands (0% descartada, 1/3 exploratoria, 2/3 media, 3/3 alta).
4. Dual piece (top + bottom) handling in Look completo with 'Look Completo' badge.
5. Consistent Candado de Afinidad Mínima naming and 0.99 universal score cap.
6. Lexical color analysis with visual embedding fallback.
7. Clean pagination avoiding awkward splits.
"""

import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def create_styled_table(doc, num_rows, col_widths, headers, data):
    total_w = sum(col_widths)
    table = doc.add_table(rows=num_rows, cols=len(col_widths))
    table.style = 'Tablaconcuadrcula'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    tblPr = table._tbl.tblPr
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_w}" w:type="dxa"/>')
    tblPr.append(tblW)
    tblLayout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
    tblPr.append(tblLayout)
    
    tblGrid = parse_xml(
        f'<w:tblGrid {nsdecls("w")}>' +
        ''.join(f'<w:gridCol w:w="{w}"/>' for w in col_widths) +
        f'</w:tblGrid>'
    )
    table._tbl.append(tblGrid)
    
    # Populate Header (Row 0)
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_width(cell, col_widths[c_idx])
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        make_run(p, title, bold=True, size_pt=9.5, color_rgb=RGBColor(255, 255, 255))
        
    # Populate Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        # Even data rows get zebra tint F5F7FA
        is_even = (r_idx % 2 == 1)
        
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_width(cell, col_widths[c_idx])
            if is_even:
                set_cell_shading(cell, "F5F7FA")
            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            lines = val.split('\n')
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(1.5)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.05
                
                format_cell_line(p, line, is_header_col=(c_idx == 0))

    return table

def format_cell_line(p, line, is_header_col=False):
    line_clean = line.strip()
    if not line_clean:
        return
        
    if is_header_col and not line_clean.startswith(('Declara:', 'Precio:', 'Catálogo:', '•')):
        make_run(p, line_clean, bold=True, size_pt=9.5)
        return
        
    if ':' in line_clean and not line_clean.startswith(('http:', 'https:')):
        parts = line_clean.split(':', 1)
        prefix = parts[0] + ':'
        rest = parts[1]
        
        make_run(p, prefix, bold=True, size_pt=9.5)
        make_run(p, rest, bold=False, size_pt=9.5)
        return
        
    if line_clean in ['Pasa', 'No pasa', 'Descartada', 'No aparece']:
        color = None
        if line_clean == 'Pasa':
            color = RGBColor(0x17, 0x36, 0x5D)
        elif line_clean in ['No pasa', 'Descartada']:
            color = RGBColor(0x71, 0x54, 0x5B)
        make_run(p, line_clean, bold=True, size_pt=9.5, color_rgb=color)
        return
        
    if any(band in line_clean for band in ['Banda Alta', 'Banda Media', 'Banda Exploratoria']):
        make_run(p, line_clean, bold=True, size_pt=9.5)
        return
        
    make_run(p, line_clean, size_pt=9.5)

def create_callout_box(doc, text_runs):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    tblPr = table._tbl.tblPr
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="10368" w:type="dxa"/>')
    tblPr.append(tblW)
    tblLayout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
    tblPr.append(tblLayout)
    
    cell = table.cell(0, 0)
    set_cell_width(cell, 10368)
    set_cell_shading(cell, "FFF2CC")
    set_cell_margins(cell, top=110, bottom=110, left=130, right=130)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    for txt, bold in text_runs:
        make_run(p, txt, bold=bold, size_pt=9.0)
        
    return table

def make_run(p, text, bold=False, italic=False, size_pt=9.5, color_rgb=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = 'Aptos'
    if size_pt:
        r.font.size = Pt(size_pt)
    if color_rgb:
        r.font.color.rgb = color_rgb
    rPr = r._r.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="es-PE"/>'))
    return r

def add_p(doc, text="", style='Normal', space_after=3, space_before=0, line_spacing=1.15, bold=False, italic=False, size_pt=None, color_rgb=None, page_break_before=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if page_break_before:
        p.paragraph_format.page_break_before = True
    if text:
        make_run(p, text, bold=bold, italic=italic, size_pt=size_pt, color_rgb=color_rgb)
    return p

def add_p_runs(doc, runs, style='Normal', space_after=3, space_before=0, line_spacing=1.15, indent_left=None, hanging=None, default_size_pt=9.0, page_break_before=False, keep_together=True, **kwargs):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if page_break_before:
        p.paragraph_format.page_break_before = True
    if keep_together:
        p.paragraph_format.keep_together = True
    
    if indent_left is not None and hanging is not None:
        pPr = p._p.get_or_add_pPr()
        ind = parse_xml(f'<w:ind {nsdecls("w")} w:left="{indent_left}" w:hanging="{hanging}"/>')
        pPr.append(ind)
        
    for r in runs:
        txt = r[0]
        bold = r[1] if len(r) > 1 else False
        italic = r[2] if len(r) > 2 else False
        size_pt = r[3] if len(r) > 3 else default_size_pt
        color_rgb = r[4] if len(r) > 4 else None
        make_run(p, txt, bold=bold, italic=italic, size_pt=size_pt, color_rgb=color_rgb)
    return p

def generate_chapter_6_elements(doc):
    elements = []
    add_b = lambda p_run: elements.append(p_run._p)
    
    # 1. Heading 1 (matched 1:1 with Chapter 7 Heading 1: 13.5pt Bold #17365D)
    p_h1 = add_p(doc, "6. Lógica del ranking: puntaje de resonancia", style='Heading 1', space_after=6, space_before=14, bold=True, size_pt=13.5, color_rgb=RGBColor(23, 54, 93))
    elements.append(p_h1._p)
    
    # 2. Callout Box (Idea Central)
    callout_runs = [
        ("Idea central: ", True),
        ("No existe un estándar externo ni un ranking fijo de «las mejores tiendas de Gamarra». Cada tienda se compara y califica ", False),
        ("únicamente contra lo que buscó la compradora en esa sesión puntual:", True),
        (" su foto de referencia, la silueta o familia de la prenda, su paleta de color y sus estilos deseados. El resultado es un puntaje relativo por tienda (de 0,0 a 1,0, presentado como porcentaje de resonancia de 0\u00A0% a 99\u00A0%) y las tiendas se ordenan de mayor a menor afinidad.", False)
    ]
    tbl_callout = create_callout_box(doc, callout_runs)
    elements.append(tbl_callout._tbl)
    
    # 3. Empty paragraph spacer
    p_sp1 = add_p(doc, "", style='Normal', space_after=3)
    elements.append(p_sp1._p)
    
    # 4. Heading 2: 6.1 Tres pasos
    p_h2_1 = add_p(doc, "6.1 Tres pasos del flujo de evaluación", style='Heading 2', space_after=4, space_before=10, bold=True, size_pt=11.5, color_rgb=RGBColor(23, 54, 93))
    elements.append(p_h2_1._p)
    
    # 5. Intro
    p_intro = add_p(doc, "El motor de búsqueda y ordenamiento de VIBRA evalúa el catálogo de Gamarra a través de tres etapas consecutivas e interconectadas:", style='Normal', space_after=4)
    elements.append(p_intro._p)
    
    # 6. Paso 1
    add_b(add_p_runs(doc, [
        ("1. Filtrar (lo exacto — Filtros duros universales). ", True),
        ("Aplica un descarte binario (sí o no) basado en la información declarada y fichas técnicas de las tiendas. Este paso es universal y aplica por igual a los tres modos de búsqueda (Una prenda, Look completo y Por estilo, salvo la aptitud de confección de la tienda, que requiere foto de referencia). Si una tienda no cumple los requisitos obligatorios de la compradora, queda excluida de inmediato antes de calcular similitudes visuales o correr modelos de inteligencia artificial:", False)
    ], indent_left=360, hanging=360, space_after=3))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Nivel de precio: ", True),
        ("Si la usuaria filtra por Económico (S/), Medio (S/S/) o Alto (S/S/S/), solo pasan las tiendas que atienden ese segmento. Si no marca ninguno, pasan todas.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Talla requerida: ", True),
        ("Si la compradora especifica una talla (XS, S, M, L, XL, Estándar), solo califican tiendas que confeccionan o comercializan dicha talla en su inventario formal.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Modalidad de compra: ", True),
        ("Si la clienta exige Solo presencial, Con envío o Venta por mayor, la condición se respeta de manera estricta.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Galería física activa: ", True),
        ("Si la usuaria selecciona una galería específica (ej. Galería El Dorado 906), el ranking se acota exclusivamente a los puestos ubicados en dicho edificio.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Aptitud de confección de la tienda (en modos con foto): ", True),
        ("La tienda debe tener en su catálogo al menos una prenda que pertenezca a la familia buscada o declarar en sus líneas que confecciona esa familia. Si un puesto se dedica exclusivamente a calzado o ropa interior y se busca un blazer, se descarta antes de puntuar.", False)
    ], indent_left=600, hanging=200, space_after=3))
    
    add_b(add_p_runs(doc, [
        ("Nota de usabilidad: ", True),
        ("Si tras aplicar los filtros duros quedan menos de 3 tiendas en el catálogo, la interfaz emite una sugerencia visible recomendando a la compradora relajar la talla o el rango de precio para ampliar sus opciones en Gamarra.", False)
    ], indent_left=600, hanging=0, space_after=5))
    
    # 7. Paso 2
    add_b(add_p_runs(doc, [
        ("2. Puntuar (lo parecido — Evaluación multidimensional). ", True),
        ("Cada tienda que superó el Paso 1 se evalúa mediante cuatro componentes técnicos, organizados en dos momentos operativos independientes:", False)
    ], indent_left=360, hanging=360, space_after=3))
    
    # Fase A: Preselección interna de prenda ganadora
    add_b(add_p_runs(doc, [
        ("Fase A: Selección de la Prenda Ganadora dentro del puesto (Filtro interno de stand). ", True),
        ("Dado que cada tienda en Gamarra tiene múltiples prendas en catálogo, el algoritmo debe determinar cuál es la pieza más representativa de esa tienda frente a lo buscado. Para cada prenda del puesto se calcula un Puntaje de Comparación interno:", False)
    ], indent_left=600, hanging=0, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("Puntaje de Comparación = Similitud Coseno × Multiplicador de Familia × Multiplicador de Color", False, True)
    ], indent_left=750, hanging=0, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Multiplicador de familia (× 1,35 o × 0,65): ", True),
        ("Si la prenda de la tienda es de la misma familia que la buscada, recibe prioridad con × 1,35. Si difiere (ej. busca blazer y la tienda ofrece pantalón), se atenúa a × 0,65. Esto garantiza que un pantalón jamás le gane a un blazer como prenda destacada, aunque compartan color.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Multiplicador de color (× 1,08): ", True),
        ("Si la prenda pertenece a la misma familia y además coincide en el grupo cromático exacto, su prioridad de selección se impulsa con × 1,08. La prenda con el mayor puntaje de comparación se convierte en la Prenda Ganadora del puesto.", False)
    ], indent_left=600, hanging=200, space_after=3))
    
    # Fase B: Medición de los 4 componentes técnicos
    add_b(add_p_runs(doc, [
        ("Fase B: Medición de los cuatro componentes de resonancia ", True),
        ("(los multiplicadores de la Fase A solo definieron qué prenda representa al stand; la resonancia final de la tienda se evalúa a partir de los siguientes componentes):", False)
    ], indent_left=600, hanging=0, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("1. Similitud visual perceptual (Gemini Embeddings 3072-d): ", True),
        ("La prenda ganadora se compara contra la foto de búsqueda calculando el coseno de sus vectores generados por gemini-embedding-2-preview. El coseno bruto se traslada mediante una curva de calibración perceptual (de ~0,55 sin relación a ~0,85 idéntico) a un Puntaje Visual limpio de 0,0 a 1,0.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("2. Clasificación en las 8 familias oficiales de confección: ", True),
        ("El catálogo se estructura taxativamente en 8 familias de prendas: ", False),
        ("1. Vestidos y enterizos", True), ("; ", False),
        ("2. Faldas y shorts", True), ("; ", False),
        ("3. Casacas, blazers y chalecos", True), (" (sastrería exterior y abrigos); ", False),
        ("4. Chompas y prendas de punto", True), ("; ", False),
        ("5. Pantalones y jeans", True), ("; ", False),
        ("6. Blusas y camisas", True), ("; ", False),
        ("7. Polos y tops", True), ("; y ", False),
        ("8. Conjuntos y sets coordinados", True), (". ", False),
        ("En conjuntos coordinados, el algoritmo inspecciona sus subpiezas (garmentContainsFamily) reconociendo si el set contiene la prenda buscada (ej. top + falda para búsqueda de falda).", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("3. Afinidad cromática léxica con respaldo visual (11 grupos cromáticos): ", True),
        ("Detección de color mediante análisis léxico con límites de palabra exacta (\\bcolor\\b) sin falsos positivos en 11 familias (Negro/Oscuro, Blanco/Marfil, Beige/Camel, Marrón/Chocolate, Azul/Denim, Rojo/Vino, Rosa/Palo Rosa, Morado/Lila, Verde/Oliva, Amarillo/Mostaza, Gris/Plomo). Si una prenda de la tienda no especifica texto cromático en su ficha, el sistema recurre al embedding visual del modelo para aproximar el tono sin aplicar bonos arbitrarios.", False)
    ], indent_left=600, hanging=200, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("4. Cobertura de estilo declarado: ", True),
        ("De los estilos que busca la compradora (máximo 3), se calcula la proporción exacta declarada formalmente por la tienda en su ficha de Gamarra: Cobertura de estilo = Estilos de la tienda coincidentes / Total de estilos buscados.", False)
    ], indent_left=600, hanging=200, space_after=4))
    
    # 8. Paso 3
    add_b(add_p_runs(doc, [
        ("3. Combinar y ordenar según el modo de búsqueda. ", True),
        ("La combinación final de los componentes depende del objetivo de compra que seleccionó la usuaria en la aplicación:", False)
    ], indent_left=360, hanging=360, space_after=4))
    
    # Table 1: Modos
    tbl1_headers = ["Modo de búsqueda", "Peso foto y silueta", "Peso estilo / cobertura", "Justificación de negocio y fórmula aplicada"]
    tbl1_widths = [2100, 1700, 2000, 4568]
    tbl1_data = [
        [
            "Una prenda",
            "Dominante\n(Silueta + Color)",
            "Modulador\n(Bono hasta +6\u00A0%)",
            "La compradora busca una pieza específica (ej. un corset marfil). Si el estilo pesara 50\u00A0%, una tienda de polos básicos que declaró «Fiesta» le ganaría a una de corsets. La silueta y el color mandan; el estilo actúa como factor de desempate."
        ],
        [
            "Look completo",
            "30\u00A0%\n(Armonía visual)",
            "70\u00A0%\n(Estilo del outfit)",
            "La compradora busca recrear una estética integral (ej. collage de Pinterest). Interesa la vibra del conjunto y la compatibilidad entre prendas superiores e inferiores, protegiendo a tiendas especialistas."
        ],
        [
            "Por estilo",
            "—",
            "Escalado por bandas\n+ Desempates",
            "Búsqueda exploratoria sin fotografía. Se ordenan directamente las tiendas cuya especialidad declarada tiene mayor intersección, complementada con desempates por estilo insignia y variedad física en catálogo."
        ]
    ]
    tbl1 = create_styled_table(doc, 4, tbl1_widths, tbl1_headers, tbl1_data)
    elements.append(tbl1._tbl)
    
    add_b(add_p(doc, "", style='Normal', space_after=2))
    
    # Detailed formulas
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Fórmula modo «Una prenda»: ", True),
        ("Resonancia = Puntaje Visual Perceptual + (Cobertura de estilo × 0,06) + Bono Color (+0,05) + Bono de alternativas (+0,02 a +0,06). El Bono Color se aplica si la prenda ganadora coincide en el grupo cromático exacto. El Bono de alternativas premia con +0,02 (hasta +0,06) por cada prenda adicional del mismo tipo con similitud coseno cruda ≥ 0,65 en el stand, asegurando que la compradora encuentre más opciones afines al llegar a la tienda física. Tope máximo: 0,99 (99\u00A0%).", False)
    ], indent_left=259, hanging=173, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Fórmula modo «Look completo»: ", True),
        ("Resonancia = (0,30 × Puntaje Visual) + (0,70 × Coincidencia de estilo de Look). ", False),
        ("Candado de Afinidad Mínima: ", True),
        ("para figurar en el ranking, la tienda debe contar obligatoriamente con al menos una prenda en catálogo con compatibilidad visual ≥ 0,45 y compartir al menos un estilo con el look; ninguna tienda califica únicamente por estilo declarado. ", False),
        ("Manejo de outfit dual y definición de especialidad: ", True),
        ("Si la tienda dispone en catálogo de prendas superiores e inferiores compatibles (ambas con afinidad ≥ 0,45), ambas se enlazan y la tienda recibe la insignia destacada de «Look Completo». La coincidencia de estilo pondera cobertura y especialidad: Coincidencia de estilo (Look) = (0,50 × Cobertura) + (0,50 × Especialidad), donde Cobertura es la proporción de estilos buscados presentes en la tienda o prendas del look, y Especialidad es la proporción de estilos declarados por la tienda que coinciden con los buscados (premiando el foco estilístico). Tope máximo: 0,99 (99\u00A0%).", False)
    ], indent_left=259, hanging=173, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Fórmula modo «Por estilo» (Con desempate multi-criterio y progresión natural de bandas): ", True),
        ("Si una tienda tiene 0 estilos coincidentes, su puntaje es 0,00 (Descartada, no aparece en los resultados). Si tiene al menos una coincidencia, la fórmula es:", False)
    ], indent_left=259, hanging=173, space_after=1))
    
    add_b(add_p_runs(doc, [
        ("Resonancia = (0,75 × Cobertura de estilo) + Bono Insignia (+0,06) + Bono de afinidad de catálogo (hasta +0,07)", True, False, 9.0, RGBColor(23, 54, 93))
    ], indent_left=450, hanging=0, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("– ", False),
        ("Escalado por bandas con 3 estilos: ", True),
        ("1 de 3 estilos (33\u00A0%) inicia en 0,25 (Banda Exploratoria); 2 de 3 estilos (67\u00A0%) inicia en 0,50 (Banda Media); y 3 de 3 estilos (100\u00A0%) alcanza 0,75 (Banda Alta). Con 3 estilos los bonos no cruzan de banda; con 2 estilos buscados, la coincidencia de 1/2 inicia en 0,375 y con bonos alcanza hasta 0,505 (pudiendo saltar de Exploratoria a Media). El denominador de la Cobertura es siempre el número de estilos seleccionados en la sesión (máximo 3); con uno solo, la escala se reduce a Descartada (0\u00A0%) o Banda Alta (75\u00A0%+).", False)
    ], indent_left=450, hanging=150, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("– ", False),
        ("Bono Insignia (+0,06) y desempate: ", True),
        ("Se activa si el estilo principal (#1 declarado) de la tienda coincide con lo buscado, elevando a una especialista sobre tiendas genéricas. Por ejemplo, una tienda especialista cuya insignia coincide y que atiende 2 de los 3 estilos solicitados parte de una base de 0,50, suma +0,06 por insignia y, con afinidad de catálogo máxima (+0,07), alcanza hasta 0,63 (Banda Media alta, a solo 0,07 del umbral de Alta); mientras que un puesto genérico con los mismos 2 estilos se queda en 0,50.", False)
    ], indent_left=450, hanging=150, space_after=2))
    
    add_b(add_p_runs(doc, [
        ("– ", False),
        ("Bono de afinidad de catálogo (hasta +0,07): ", True),
        ("Premia la afinidad promedio de las prendas de la tienda con los estilos buscados. Es un valor continuo que desempata naturalmente tiendas con igual coincidencia declarada. Tope máximo universal: 0,99 (99\u00A0%).", False)
    ], indent_left=450, hanging=150, space_after=3))
    
    add_b(add_p_runs(doc, [
        ("• ", False),
        ("Factores fuera del puntaje: ", True),
        ("La fecha de actualización se exhibe visiblemente en la tarjeta, la ubicación en galería se utiliza para ordenar el recorrido de compra en el mapa y el feedback de la usuaria se almacena para calibración continua.", False)
    ], indent_left=259, hanging=173, space_after=5))
    
    # 9. Heading 2: 6.2 Bandas (with clean page break before to keep table on one page!)
    p_h2_2 = add_p(doc, "6.2 Cómo se presenta el resultado: bandas de resonancia", style='Heading 2', space_after=4, space_before=10, bold=True, size_pt=11.5, color_rgb=RGBColor(23, 54, 93), page_break_before=True)
    elements.append(p_h2_2._p)
    
    # Table 2: Bandas
    tbl2_headers = ["Banda de resonancia", "Puntaje de Resonancia", "Presentación visual y significado para la compradora"]
    tbl2_widths = [2000, 2200, 6168]
    tbl2_data = [
        [
            "Alta",
            "70\u00A0% o más (≥ 0,70)",
            "Tres arcos completos del medidor en color Higo 70 (#71545B). Coincidencia sobresaliente en silueta y estética de estilo requerida (con afinidad cromática coincidente si aplica en modo de prenda individual)."
        ],
        [
            "Media",
            "45\u00A0% a 69\u00A0% (0,45 a 0,69)",
            "Dos arcos activos en color Grafito 55 (#676B63). Tiendas con prendas similares o del mismo estilo, con ligeras variaciones de corte o confección."
        ],
        [
            "Exploratoria",
            "25\u00A0% a 44\u00A0% (0,25 a 0,44)",
            "Un arco activo en color Grafito 55 (#676B63). Tiendas que trabajan la categoría o estética pero con siluetas alternativas para la clienta que desea explorar."
        ],
        [
            "No se muestra",
            "Menos de 25\u00A0% (< 0,25)",
            "La tienda no califica para el ranking de esa búsqueda. Si ninguna tienda supera el umbral, la app muestra un mensaje claro orientativo: «Ninguna tienda superó el umbral de afinidad con esta foto. Prueba con un encuadre más cerrado de la prenda o quita los filtros de precio y talla»."
        ]
    ]
    tbl2 = create_styled_table(doc, 5, tbl2_widths, tbl2_headers, tbl2_data)
    elements.append(tbl2._tbl)
    
    add_b(add_p(doc, "", style='Normal', space_after=2))
    
    add_b(add_p_runs(doc, [
        ("Regla de asignación y redondeo: ", True),
        ("El puntaje numérico final se redondea a dos decimales estándar antes de catalogarse en la banda correspondiente (por ejemplo, 0,695 redondea a 0,70, ingresando a Banda Alta). Las bandas de resonancia no comparan a las tiendas contra un estándar abstracto ni califican su calidad comercial intrínseca: ", False),
        ("solo expresan con transparencia el grado de alineación entre el inventario de la tienda y la búsqueda puntual de la compradora.", True),
        (" La resonancia se comunica siempre como una recomendación personalizada de compra, nunca como una garantía absoluta de disponibilidad física inmediata.", False)
    ], space_after=5))
    
    # 10. Heading 2: 6.3 Ejemplo (with clean page break before to keep table on one page!)
    p_h2_3 = add_p(doc, "6.3 Ejemplo práctico paso a paso (Caso real en Gamarra)", style='Heading 2', space_after=4, space_before=10, bold=True, size_pt=11.5, color_rgb=RGBColor(23, 54, 93), page_break_before=True)
    elements.append(p_h2_3._p)
    
    add_b(add_p_runs(doc, [
        ("Escenario de búsqueda real: ", True),
        ("Una compradora sube la foto de un conjunto sastre de ", False),
        ("chaleco marfil con short a juego", True),
        (" (perteneciente a la familia de casacas, blazers y chalecos) en modo ", False),
        ("Look completo", True),
        (". Selecciona filtro de precio ", False),
        ("Medio (S/S/)", True),
        (" y confirma los estilos ", False),
        ("Elegante", True),
        (" y ", False),
        ("Oficina", True),
        (".", False)
    ], space_after=3))
    
    add_b(add_p_runs(doc, [
        ("Nota de procesamiento visual: ", True),
        ("El modelo multimodal de Gemini Embeddings procesa la imagen global del outfit capturando su silueta y paleta cromática; las prendas físicas de cada stand se evalúan contra este vector representativo identificando la prenda protagónica y piezas complementarias sin requerir segmentación asistida. Los puntajes visuales calibrados corresponden a la función de calibración perceptual del motor de ranking (0,82 → 0,86 y 0,60 → 0,20).", False)
    ], space_after=4))
    
    # Table 3: Ejemplo
    tbl3_headers = ["Tienda evaluada", "Paso 1: Filtro exacto", "Paso 2: Análisis técnico", "Paso 3: Ponderación Look\n(30\u00A0% Foto + 70\u00A0% Estilo)", "Resultado y Banda"]
    tbl3_widths = [2500, 1600, 2200, 2168, 1900]
    tbl3_data = [
        [
            "Tienda A (Beige Studio)\nDeclara: Elegante, Oficina\nPrecio: Medio (S/S/)\nCatálogo: Chalecos sastre y palazzos marfil",
            "Pasa\n(Cumple precio medio)",
            "• Prenda ganadora: Chaleco sastre\n• Coseno crudo: 0,82\n• Puntaje visual calibrado: 0,86\n• Estilos: 2/2 = 1,00 (cobertura 1,0; especialidad 1,0)\n• Arma outfit: Sí (chaleco + palazzo)",
            "(0,30 × 0,86) + (0,70 × 1,00)\n= 0,258 + 0,70\n= 0,958 ≈ 0,96",
            "96\u00A0% · Banda Alta\n(Tres arcos Higo)\nPuesto #1 del ranking\nInsignia «Look Completo»"
        ],
        [
            "Tienda B (Urbana Chic)\nDeclara: Casual, Urbano\nPrecio: Medio (S/S/)\nCatálogo: Shorts de lino, sin chaleco sastre",
            "Pasa\n(Cumple precio medio)",
            "• Prenda ganadora: Short lino\n• Coseno crudo: 0,60\n• Calibrado: 0,20\n• Cobertura de estilo: 0/2 = 0,00\n• Arma outfit: No compatible\n• Candado de Afinidad Mínima: No supera (0,20 < 0,45)",
            "(0,30 × 0,20) + (0,70 × 0,00)\n= 0,060 + 0,00\n= 0,06",
            "Descartada\n(Puntaje 6\u00A0% < 25\u00A0% y no supera Candado Mínimo; no se muestra en el ranking)"
        ],
        [
            "Tienda C (Alta Costura Gamarra)\nDeclara: Elegante, Oficina\nPrecio: Alto (S/S/S/)\nCatálogo: Conjuntos sastre idénticos",
            "No pasa\n(Precio alto ≠ precio medio seleccionado)",
            "No se calcula\n(Descarte temprano en Paso 1)",
            "—",
            "No aparece\n(Filtrada en Paso 1, ahorrando cómputo)"
        ]
    ]
    tbl3 = create_styled_table(doc, 4, tbl3_widths, tbl3_headers, tbl3_data)
    elements.append(tbl3._tbl)
    
    add_b(add_p(doc, "", style='Normal', space_after=2))
    
    add_b(add_p_runs(doc, [
        ("Análisis del resultado: ", True),
        ("Tienda A (Beige Studio) obtiene el puesto #1 con 96\u00A0% porque satisface tanto la silueta física del chaleco sastre como la intención completa de estilo, además de ofrecer prendas combinables para armar el look completo; Tienda B (Urbana Chic) es descartada del ranking porque su inventario no corresponde a la silueta sastre requerida (lo que reduce su similitud visual calibrada a 0,20) ni coincide con los estilos solicitados, evitando frustración a la clienta; y Tienda C (Alta Costura Gamarra) es filtrada de inmediato en el Paso 1 por tener precio Alto, respetando estrictamente el presupuesto definido por la compradora.", False)
    ], space_after=5))
    
    # 11. Heading 2: 6.4 Comparativa
    p_h2_4 = add_p(doc, "6.4 Comparativa: Propuesta inicial vs. Implementación real calibrada", style='Heading 2', space_after=4, space_before=10, bold=True, size_pt=11.5, color_rgb=RGBColor(23, 54, 93), page_break_before=True)
    elements.append(p_h2_4._p)
    
    add_b(add_p(doc, "Esta tabla sintetiza las decisiones de ingeniería de producto que diferenciaron la hipótesis teórica inicial de la implementación en producción, validadas en la calibración inicial con las 23 tiendas y 277 prendas reales de Gamarra:", style='Normal', space_after=4))
    
    # Table 4: Comparativa
    tbl4_headers = ["Aspecto evaluado", "Propuesta teórica inicial", "Implementación calibrada en producción", "Justificación y hallazgo empírico en Gamarra"]
    tbl4_widths = [2000, 2100, 2700, 3568]
    tbl4_data = [
        [
            "Modo Una prenda",
            "50\u00A0% Foto + 50\u00A0% Estilo",
            "Silueta dominante + Color + Estilo modulador (+6\u00A0%)",
            "Si el estilo pesaba 50\u00A0%, tiendas de polos básicos con etiqueta «Fiesta» superaban a tiendas de vestidos. La prenda física y su silueta deben mandar siempre."
        ],
        [
            "Modo Look completo",
            "30\u00A0% Foto + 70\u00A0% Estilo",
            "30\u00A0% Foto + 70\u00A0% Estilo + Candado de Afinidad Mínima",
            "Se mantuvo la proporción base 30/70 exigiendo el Candado de Afinidad Mínima (visual ≥ 0,45 y al menos 1 estilo compartido), lo que evita recomendar puestos sin relación visual. Reconoce que en outfit completo la coherencia de estilo manda (70\u00A0%), mientras que el filtro visual mínimo de 0,45 previene desalineaciones estéticas evidentes entre lo buscado y lo exhibido."
        ],
        [
            "Modo Por estilo",
            "100\u00A0% Coincidencia directa",
            "Escala por bandas (0,75) + Bono Insignia + Bono Afinidad Catálogo",
            "La proporción directa provocaba empates masivos (33\u00A0%, 67\u00A0%, 100\u00A0%). El algoritmo introduce desempates por estilo insignia y afinidad de prendas en catálogo físico."
        ],
        [
            "Discriminación de color",
            "Similitud visual general de la foto",
            "Reconocimiento por 11 grupos cromáticos con límites léxicos",
            "La similitud vectorial pura a veces confundía texturas idénticas de diferente color. El filtro léxico garantiza que el color coincida y otorga bono explícito (+0,05)."
        ],
        [
            "Conjuntos de 2 piezas",
            "Tratados como prendas individuales",
            "Desglose de subpiezas (garmentContainsFamily)",
            "En Gamarra abundan los sets de 2 o 3 piezas (top + falda, chaleco + short). El sistema ahora reconoce si un conjunto contiene la subprenda que la clienta busca."
        ],
        [
            "Umbrales de bandas",
            "85\u00A0% Alta / 70\u00A0% Media / 55\u00A0% Exploratoria",
            "70\u00A0% Alta / 45\u00A0% Media / 25\u00A0% Exploratoria",
            "La distribución natural de similitud textil en embeddings requiere umbrales calibrados para que la Banda Alta sea selectiva pero alcanzable por las mejores tiendas."
        ],
        [
            "Bono de alternativas en stand",
            "No contemplado",
            "+0,02 a +0,06 por opciones adicionales afines en tienda",
            "En la experiencia de compra en Gamarra, llegar a un stand con 3 o 4 alternativas similares eleva sustancialmente la satisfacción y conversión de la compradora."
        ]
    ]
    tbl4 = create_styled_table(doc, 8, tbl4_widths, tbl4_headers, tbl4_data)
    elements.append(tbl4._tbl)
    
    add_b(add_p(doc, "", style='Normal', space_after=2))
    
    # 12. Final Technical Note
    add_b(add_p_runs(doc, [
        ("Nota técnica de calibración y MLOps: ", True),
        ("El modelo de embeddings opera en un espacio vectorial de 3072 dimensiones bajo el modelo gemini-embedding-2-preview. Dos prendas sin relación promedian una similitud coseno de ~0,55 y dos prendas muy similares alcanzan ~0,85. La función de normalización perceptual implementada mapea dicho intervalo de forma no lineal a la escala 0\u00A0% – 99\u00A0%, asegurando que los porcentajes reflejen con fidelidad la percepción visual humana. Asimismo, como buena práctica de arquitectura de datos, cada vector generado almacena el identificador y versión exacta del modelo para garantizar reproducibilidad en futuras expansiones del catálogo.", False)
    ], space_after=8, default_size_pt=9.0))
    
    return elements

def run_update(input_path, output_path):
    print(f"Loading document: {input_path}")
    doc = docx.Document(input_path)
    body = doc._body._body
    
    start_idx = None
    end_idx = None
    
    for idx, child in enumerate(body):
        if isinstance(child, CT_P):
            p = docx.text.paragraph.Paragraph(child, doc)
            if p.text.startswith(('6. Lógica', '6. L\xf3gica')):
                start_idx = idx
            elif p.text.startswith('7. Base de'):
                end_idx = idx
                break
                
    if start_idx is None or end_idx is None:
        raise ValueError(f"Could not locate Chapter 6 boundaries (start={start_idx}, end={end_idx})")
        
    print(f"Located Chapter 6 from body index {start_idx} to {end_idx} (exclusive)")
    
    old_elements = [body[i] for i in range(start_idx, end_idx)]
    end_element = body[end_idx]
    
    print("Generating refined Chapter 6 elements...")
    new_elements = generate_chapter_6_elements(doc)
    print(f"Generated {len(new_elements)} refined elements.")
    
    for elem in new_elements:
        body.remove(elem)
        end_element.addprevious(elem)
        
    print(f"Removing {len(old_elements)} old Chapter 6 elements...")
    for elem in old_elements:
        body.remove(elem)
        
    # Ensure Chapter 7 starts cleanly on a fresh page
    p_ch7 = docx.text.paragraph.Paragraph(end_element, doc)
    p_ch7.paragraph_format.page_break_before = True
    print("Set page_break_before=True on Chapter 7 heading.")
    
    # Ensure all tables across the entire document (including Table 7.1) have cantSplit on rows
    for tbl in doc.tables:
        for row in tbl.rows:
            trPr = row._tr.get_or_add_trPr()
            if not row._tr.xpath("./w:trPr/w:cantSplit"):
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    print("Enforced cantSplit on all document table rows.")

    # Enforce es-PE on all document styles and all paragraphs/runs to eliminate Word spelling squiggles
    for s in doc.styles:
        if hasattr(s, '_element') and s._element is not None:
            rPr = s._element.get_or_add_rPr()
            if not rPr.xpath("./w:lang"):
                rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="es-PE"/>'))
                
    for p in doc.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            if not rPr.xpath("./w:lang"):
                rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="es-PE"/>'))
                
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        rPr = r._r.get_or_add_rPr()
                        if not rPr.xpath("./w:lang"):
                            rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="es-PE"/>'))
    print("Enforced es-PE language attribute across all document styles and text runs.")
        
    print(f"Saving updated document to: {output_path}")
    doc.save(output_path)
    print("Done!")

if __name__ == '__main__':
    src = "c:/Proyectos/VIBRA/VIBRA.docx"
    if not os.path.exists(src):
        src = "c:/Proyectos/VIBRA/VIBRA.backup.docx"
    
    # Try primary outputs, falling back gracefully if Word has one locked
    output_candidates = [
        "c:/Proyectos/VIBRA/VIBRA.docx",
        "c:/Proyectos/VIBRA/VIBRA.final.docx",
        "c:/Proyectos/VIBRA/VIBRA.actualizado.docx",
        "c:/Proyectos/VIBRA/VIBRA.v2.docx"
    ]
    
    saved_paths = []
    # First save to a temp or available candidate
    temp_target = "c:/Proyectos/VIBRA/VIBRA.temp.docx"
    run_update(src, temp_target)
    
    import shutil
    for path in output_candidates:
        try:
            shutil.copy2(temp_target, path)
            saved_paths.append(path)
            print(f"Successfully saved to: {path}")
        except Exception as e:
            print(f"Could not write to {path} (likely open in Word): {e}")
            
    if os.path.exists(temp_target):
        try:
            os.remove(temp_target)
        except Exception:
            pass
            
    print(f"Finished! Successfully updated: {saved_paths}")
